from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from flask_jwt_extended import JWTManager, create_access_token, jwt_required, get_jwt_identity
from flask_sqlalchemy import SQLAlchemy
from flask_socketio import SocketIO, emit
from flask_mail import Mail, Message
from functools import wraps
from datetime import datetime, timedelta
from werkzeug.utils import secure_filename
from sqlalchemy import func, or_
from faker import Faker
import bcrypt, random, os, mimetypes

# ==================== APP SETUP ====================
app = Flask(__name__)
CORS(app, resources={
    r"/api/*": {
        "origins": "*",
        "methods": ["GET", "POST", "PUT", "DELETE", "OPTIONS"],
        "allow_headers": ["Content-Type", "Authorization"]
    }
})

app.config['SECRET_KEY'] = 'freightpulse-secret-key-2026-very-long-and-secure-string-xyz'
app.config['JWT_SECRET_KEY'] = 'freightpulse-jwt-secret-key-2026-very-long-and-secure-string-abc'
app.config['JWT_ACCESS_TOKEN_EXPIRES'] = timedelta(hours=8)
app.config['SQLALCHEMY_DATABASE_URI'] = 'postgresql://postgres:Angel%405326@localhost/freight_db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

UPLOAD_FOLDER = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'uploads')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['MAX_CONTENT_LENGTH'] = 20 * 1024 * 1024

app.config['MAIL_SERVER'] = 'smtp.gmail.com'
app.config['MAIL_PORT'] = 587
app.config['MAIL_USE_TLS'] = True
app.config['MAIL_USERNAME'] = os.environ.get('MAIL_USER', 'your-email@gmail.com')
app.config['MAIL_PASSWORD'] = os.environ.get('MAIL_PASS', 'your-app-password')
app.config['MAIL_DEFAULT_SENDER'] = app.config['MAIL_USERNAME']

db = SQLAlchemy(app)
jwt = JWTManager(app)
socketio = SocketIO(app, cors_allowed_origins="*")
mail = Mail(app)
fake = Faker()

# ==================== MODELS ====================
class Organization(db.Model):
    __tablename__ = 'organizations'
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(150), nullable=False)
    slug = db.Column(db.String(50), unique=True, nullable=False)
    plan = db.Column(db.String(20), default='free')
    is_active = db.Column(db.Boolean, default=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


class User(db.Model):
    __tablename__ = 'users'
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    password_hash = db.Column(db.String(128), nullable=False)
    role = db.Column(db.String(20), default='user')
    email = db.Column(db.String(120))
    phone = db.Column(db.String(20))
    department = db.Column(db.String(50))
    organization_id = db.Column(db.Integer, db.ForeignKey('organizations.id'))
    is_active = db.Column(db.Boolean, default=True)
    last_login = db.Column(db.DateTime)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

    def set_password(self, p):
        self.password_hash = bcrypt.hashpw(p.encode(), bcrypt.gensalt()).decode()

    def check_password(self, p):
        return bcrypt.checkpw(p.encode(), self.password_hash.encode())


class Client(db.Model):
    __tablename__ = 'clients'
    id = db.Column(db.Integer, primary_key=True)
    organization_id = db.Column(db.Integer, db.ForeignKey('organizations.id'))
    company_name = db.Column(db.String(150), nullable=False)
    contact_person = db.Column(db.String(100))
    email = db.Column(db.String(120))
    phone = db.Column(db.String(20))
    address = db.Column(db.String(200))
    city = db.Column(db.String(80))
    country = db.Column(db.String(80))
    tax_id = db.Column(db.String(50))
    credit_limit = db.Column(db.Float, default=50000)
    is_active = db.Column(db.Boolean, default=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


class Shipment(db.Model):
    __tablename__ = 'shipments'
    id = db.Column(db.Integer, primary_key=True)
    organization_id = db.Column(db.Integer, db.ForeignKey('organizations.id'))
    shipment_number = db.Column(db.String(30), unique=True, nullable=False)
    client_id = db.Column(db.Integer, db.ForeignKey('clients.id'))
    client_name = db.Column(db.String(150))
    origin = db.Column(db.String(100))
    destination = db.Column(db.String(100))
    status = db.Column(db.String(30), default='Booked')
    carrier = db.Column(db.String(100))
    vessel_flight = db.Column(db.String(100))
    etd = db.Column(db.DateTime)
    eta = db.Column(db.DateTime)
    currency = db.Column(db.String(5), default='USD')
    total_value = db.Column(db.Float, default=0)
    payment_status = db.Column(db.String(20), default='Unpaid')
    is_deleted = db.Column(db.Boolean, default=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    latitude = db.Column(db.Float)
    longitude = db.Column(db.Float)
    items = db.relationship('ShipmentItem', backref='shipment', cascade='all, delete-orphan')
    events = db.relationship('TrackingEvent', backref='shipment', cascade='all, delete-orphan', order_by='TrackingEvent.created_at')
    documents = db.relationship('Document', backref='shipment', cascade='all, delete-orphan')


class ShipmentItem(db.Model):
    __tablename__ = 'shipment_items'
    id = db.Column(db.Integer, primary_key=True)
    shipment_id = db.Column(db.Integer, db.ForeignKey('shipments.id'), nullable=False)
    description = db.Column(db.String(200), nullable=False)
    quantity = db.Column(db.Integer, default=1)
    weight_kg = db.Column(db.Float, default=0)
    declared_value = db.Column(db.Float, default=0)
    hs_code = db.Column(db.String(20))
    package_type = db.Column(db.String(50))


class TrackingEvent(db.Model):
    __tablename__ = 'tracking_events'
    id = db.Column(db.Integer, primary_key=True)
    shipment_id = db.Column(db.Integer, db.ForeignKey('shipments.id'), nullable=False)
    status = db.Column(db.String(30), nullable=False)
    location = db.Column(db.String(100))
    notes = db.Column(db.Text)
    created_by = db.Column(db.String(80))
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


class Invoice(db.Model):
    __tablename__ = 'invoices'
    id = db.Column(db.Integer, primary_key=True)
    organization_id = db.Column(db.Integer, db.ForeignKey('organizations.id'))
    invoice_number = db.Column(db.String(30), unique=True, nullable=False)
    shipment_id = db.Column(db.Integer, db.ForeignKey('shipments.id'))
    client_name = db.Column(db.String(150))
    amount = db.Column(db.Float, default=0)
    tax = db.Column(db.Float, default=0)
    total = db.Column(db.Float, default=0)
    paid_amount = db.Column(db.Float, default=0)
    currency = db.Column(db.String(5), default='USD')
    status = db.Column(db.String(20), default='Unpaid')
    issue_date = db.Column(db.DateTime, default=datetime.utcnow)
    due_date = db.Column(db.DateTime)
    paid_date = db.Column(db.DateTime)
    notes = db.Column(db.Text)
    payments = db.relationship('Payment', backref='invoice', cascade='all, delete-orphan')

    @property
    def balance(self):
        return round((self.total or 0) - (self.paid_amount or 0), 2)


class Payment(db.Model):
    __tablename__ = 'payments'
    id = db.Column(db.Integer, primary_key=True)
    invoice_id = db.Column(db.Integer, db.ForeignKey('invoices.id'), nullable=False)
    amount = db.Column(db.Float, default=0)
    method = db.Column(db.String(30), default='Bank Transfer')
    reference = db.Column(db.String(100))
    notes = db.Column(db.Text)
    paid_date = db.Column(db.DateTime, default=datetime.utcnow)
    recorded_by = db.Column(db.String(80))


class Document(db.Model):
    __tablename__ = 'documents'
    id = db.Column(db.Integer, primary_key=True)
    shipment_id = db.Column(db.Integer, db.ForeignKey('shipments.id'), nullable=False)
    filename = db.Column(db.String(200), nullable=False)
    original_name = db.Column(db.String(200))
    file_type = db.Column(db.String(50))
    file_size = db.Column(db.Integer, default=0)
    mime_type = db.Column(db.String(100))
    uploaded_by = db.Column(db.String(80))
    uploaded_at = db.Column(db.DateTime, default=datetime.utcnow)


class Notification(db.Model):
    __tablename__ = 'notifications'
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer)
    title = db.Column(db.String(200))
    message = db.Column(db.Text)
    type = db.Column(db.String(30), default='info')
    is_read = db.Column(db.Boolean, default=False)
    link = db.Column(db.String(200))
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


class ActivityLog(db.Model):
    __tablename__ = 'activity_logs'
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer)
    username = db.Column(db.String(80))
    action = db.Column(db.String(50))
    entity = db.Column(db.String(50))
    entity_id = db.Column(db.Integer)
    details = db.Column(db.Text)
    ip_address = db.Column(db.String(50))
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


# ==================== HELPERS ====================
def log_activity(action, entity, entity_id, details='', user_id=None, username=None):
    """Log an activity. Can accept user_id/username directly or pull from JWT context."""
    try:
        if user_id is None:
            try:
                uid = get_jwt_identity()
                user_id = int(uid) if uid else None
            except Exception:
                user_id = None

        if username is None and user_id:
            u = db.session.get(User, user_id)
            username = u.username if u else 'system'

        log = ActivityLog(
            user_id=user_id,
            username=username or 'system',
            action=action, entity=entity, entity_id=entity_id,
            details=details,
            ip_address=request.remote_addr if request else None
        )
        db.session.add(log)
        db.session.commit()
    except Exception as e:
        print(f"Log error: {e}")


def notify(user_id, title, message, type='info', link=''):
    try:
        n = Notification(user_id=user_id, title=title, message=message, type=type, link=link)
        db.session.add(n)
        db.session.commit()
        socketio.emit('new_notification', {'title': title, 'message': message, 'type': type}, broadcast=True)
    except Exception as e:
        print(f"Notify error: {e}")


def send_email(to, subject, html_body):
    try:
        msg = Message(subject=subject, recipients=[to], html=html_body)
        mail.send(msg)
        return True
    except Exception as e:
        print(f"Email skipped: {e}")
        return False


def shipment_email_html(shipment, event):
    return f"""
    <div style="font-family: Arial, sans-serif; background:#0a1628; padding:30px; color:white;">
      <div style="max-width:600px; margin:auto; background:#1a2a4a; border-radius:16px; padding:30px; border:1px solid #C9A84C;">
        <h1 style="color:#C9A84C; margin:0;">FreightPulse</h1>
        <h2 style="color:white;">{shipment.shipment_number}</h2>
        <p>Status: <strong style="color:#C9A84C;">{event.status}</strong></p>
        <p>Location: {event.location or '—'}</p>
        <p>{event.notes or ''}</p>
      </div>
    </div>
    """


def admin_required():
    def wrapper(fn):
        @wraps(fn)
        def deco(*a, **kw):
            uid = get_jwt_identity()
            u = db.session.get(User, int(uid)) if uid else None
            if not u or u.role != 'admin':
                return jsonify({'msg': 'Admin required'}), 403
            return fn(*a, **kw)
        return deco
    return wrapper


def get_current_org():
    try:
        uid = get_jwt_identity()
        if not uid:
            return None
        user = db.session.get(User, int(uid))
        return user.organization_id if user else None
    except Exception:
        return None


ALLOWED_EXT = {'pdf', 'png', 'jpg', 'jpeg', 'doc', 'docx', 'xls', 'xlsx', 'txt'}


def allowed_file(name):
    return '.' in name and name.rsplit('.', 1)[1].lower() in ALLOWED_EXT


# ==================== INIT DB ====================
with app.app_context():
    db.create_all()

    if not Organization.query.first():
        org = Organization(name='FreightPulse Demo Co.', slug='demo', plan='enterprise')
        db.session.add(org)
        db.session.commit()
        print("✅ Default organization created")

    default_org = Organization.query.first()

    if not User.query.filter_by(username='admin').first():
        a = User(
            username='admin', role='admin',
            email='admin@freightpulse.com',
            organization_id=default_org.id
        )
        a.set_password('admin123')
        db.session.add(a)
        db.session.commit()
        print("✅ Admin created: admin/admin123")


# ==================== AUTH ====================
@app.route('/')
def home():
    return "🚀 FreightPulse v5.0"


@app.route('/api/auth/signup', methods=['POST'])
def signup():
    d = request.get_json()
    if User.query.filter_by(username=d.get('username')).first():
        return jsonify({'msg': 'Username taken'}), 409

    org_id = d.get('organization_id')
    if not org_id:
        org_name = d.get('organization_name', f"{d['username']}'s Org")
        slug = d['username'].lower().replace(' ', '-')[:40]
        base = slug
        i = 1
        while Organization.query.filter_by(slug=slug).first():
            slug = f"{base}-{i}"
            i += 1
        org = Organization(name=org_name, slug=slug, plan='free')
        db.session.add(org)
        db.session.flush()
        org_id = org.id

    u = User(
        username=d['username'],
        email=d.get('email', ''),
        role='admin',
        organization_id=org_id
    )
    u.set_password(d['password'])
    db.session.add(u)
    db.session.commit()
    return jsonify({'msg': 'Created!', 'organization_id': org_id}), 201


@app.route('/api/auth/login', methods=['POST'])
def login():
    d = request.get_json()
    u = User.query.filter_by(username=d.get('username')).first()
    if not u or not u.check_password(d.get('password', '')):
        return jsonify({'msg': 'Invalid credentials'}), 401
    if not u.is_active:
        return jsonify({'msg': 'Account disabled'}), 403

    u.last_login = datetime.utcnow()
    db.session.commit()

    t = create_access_token(identity=str(u.id))
    # ✅ Pass user info directly (login isn't JWT-protected)
    log_activity('LOGIN', 'user', u.id, 'Login successful',
                 user_id=u.id, username=u.username)

    return jsonify({
        'data': {
            'access_token': t,
            'user': {
                'id': u.id, 'username': u.username, 'role': u.role,
                'email': u.email, 'organization_id': u.organization_id
            }
        }
    }), 200


# ==================== CLIENTS ====================
@app.route('/api/clients/', methods=['GET'])
@jwt_required()
def list_clients():
    org_id = get_current_org()
    q = request.args.get('q', '').lower()
    clients = Client.query.filter_by(organization_id=org_id)
    if q:
        clients = clients.filter(Client.company_name.ilike(f'%{q}%'))
    clients = clients.order_by(Client.id.desc()).all()
    return jsonify([{
        'id': c.id, 'company_name': c.company_name,
        'contact_person': c.contact_person, 'email': c.email,
        'phone': c.phone, 'address': c.address, 'city': c.city,
        'country': c.country, 'tax_id': c.tax_id,
        'credit_limit': c.credit_limit, 'is_active': c.is_active,
        'created_at': c.created_at.isoformat() if c.created_at else None
    } for c in clients]), 200


@app.route('/api/clients/', methods=['POST'])
@jwt_required()
def create_client():
    d = request.get_json()
    c = Client(
        organization_id=get_current_org(),
        **{k: d.get(k) for k in ['company_name', 'contact_person', 'email',
                                  'phone', 'address', 'city', 'country', 'tax_id']}
    )
    c.credit_limit = float(d.get('credit_limit', 50000))
    db.session.add(c)
    db.session.commit()
    log_activity('CREATE', 'client', c.id, c.company_name)
    return jsonify({'msg': 'Client added', 'id': c.id}), 201


@app.route('/api/clients/<int:cid>', methods=['PUT'])
@jwt_required()
def update_client(cid):
    c = db.session.get(Client, cid)
    if not c:
        return jsonify({'msg': 'Not found'}), 404
    d = request.get_json()
    for k in ['company_name', 'contact_person', 'email', 'phone',
              'address', 'city', 'country', 'tax_id']:
        if k in d:
            setattr(c, k, d[k])
    if 'credit_limit' in d:
        c.credit_limit = float(d['credit_limit'])
    if 'is_active' in d:
        c.is_active = d['is_active']
    db.session.commit()
    log_activity('UPDATE', 'client', cid, c.company_name)
    return jsonify({'msg': 'Updated'}), 200


@app.route('/api/clients/<int:cid>', methods=['DELETE'])
@jwt_required()
def delete_client(cid):
    c = db.session.get(Client, cid)
    if not c:
        return jsonify({'msg': 'Not found'}), 404
    db.session.delete(c)
    db.session.commit()
    log_activity('DELETE', 'client', cid, c.company_name)
    return jsonify({'msg': 'Deleted'}), 200


@app.route('/api/clients/<int:cid>/balance', methods=['GET'])
@jwt_required()
def client_balance(cid):
    c = db.session.get(Client, cid)
    if not c:
        return jsonify({'msg': 'Not found'}), 404
    shipments = Shipment.query.filter_by(client_id=cid, is_deleted=False).all()
    ids = [s.id for s in shipments]
    invoices = Invoice.query.filter(Invoice.shipment_id.in_(ids)).all() if ids else []
    ti = sum(i.total or 0 for i in invoices)
    tp = sum(i.paid_amount or 0 for i in invoices)
    return jsonify({
        'client_name': c.company_name,
        'credit_limit': c.credit_limit or 0,
        'total_invoiced': round(ti, 2),
        'total_paid': round(tp, 2),
        'outstanding': round(ti - tp, 2),
        'credit_available': round((c.credit_limit or 0) - (ti - tp), 2)
    }), 200


# ==================== SHIPMENTS ====================
@app.route('/api/shipments/', methods=['GET'])
@jwt_required()
def list_shipments():
    org_id = get_current_org()

    status = request.args.get('status')
    client_id = request.args.get('client_id')
    search = request.args.get('q', '').strip()
    from_date = request.args.get('from_date')
    to_date = request.args.get('to_date')
    min_value = request.args.get('min_value', type=float)
    max_value = request.args.get('max_value', type=float)
    origin = request.args.get('origin', '').strip()
    destination = request.args.get('destination', '').strip()
    carrier = request.args.get('carrier', '').strip()

    q = Shipment.query.filter_by(is_deleted=False, organization_id=org_id)

    if status and status != 'All':
        q = q.filter_by(status=status)
    if client_id:
        q = q.filter_by(client_id=int(client_id))
    if search:
        like = f'%{search}%'
        q = q.filter(or_(
            Shipment.shipment_number.ilike(like),
            Shipment.client_name.ilike(like),
            Shipment.origin.ilike(like),
            Shipment.destination.ilike(like),
            Shipment.carrier.ilike(like)
        ))
    if from_date:
        q = q.filter(Shipment.created_at >= from_date)
    if to_date:
        q = q.filter(Shipment.created_at <= to_date + ' 23:59:59')
    if min_value is not None:
        q = q.filter(Shipment.total_value >= min_value)
    if max_value is not None:
        q = q.filter(Shipment.total_value <= max_value)
    if origin:
        q = q.filter(Shipment.origin.ilike(f'%{origin}%'))
    if destination:
        q = q.filter(Shipment.destination.ilike(f'%{destination}%'))
    if carrier:
        q = q.filter(Shipment.carrier.ilike(f'%{carrier}%'))

    shipments = q.order_by(Shipment.id.desc()).all()
    return jsonify([{
        'id': s.id, 'shipment_number': s.shipment_number,
        'client_id': s.client_id, 'client_name': s.client_name,
        'origin': s.origin, 'destination': s.destination, 'status': s.status,
        'carrier': s.carrier, 'vessel_flight': s.vessel_flight,
        'etd': s.etd.isoformat() if s.etd else None,
        'eta': s.eta.isoformat() if s.eta else None,
        'currency': s.currency, 'total_value': s.total_value,
        'payment_status': s.payment_status,
        'latitude': s.latitude, 'longitude': s.longitude,
        'created_at': s.created_at.isoformat() if s.created_at else None,
        'items_count': len(s.items), 'events_count': len(s.events),
        'documents_count': len(s.documents)
    } for s in shipments]), 200


@app.route('/api/shipments/<int:sid>', methods=['GET'])
@jwt_required()
def get_shipment(sid):
    s = db.session.get(Shipment, sid)
    if not s or s.is_deleted:
        return jsonify({'msg': 'Not found'}), 404
    return jsonify({
        'id': s.id, 'shipment_number': s.shipment_number,
        'client_id': s.client_id, 'client_name': s.client_name,
        'origin': s.origin, 'destination': s.destination, 'status': s.status,
        'carrier': s.carrier, 'vessel_flight': s.vessel_flight,
        'etd': s.etd.isoformat() if s.etd else None,
        'eta': s.eta.isoformat() if s.eta else None,
        'currency': s.currency, 'total_value': s.total_value,
        'payment_status': s.payment_status,
        'latitude': s.latitude, 'longitude': s.longitude,
        'created_at': s.created_at.isoformat() if s.created_at else None,
        'items': [{'id': i.id, 'description': i.description,
                   'quantity': i.quantity, 'weight_kg': i.weight_kg,
                   'declared_value': i.declared_value,
                   'hs_code': i.hs_code, 'package_type': i.package_type}
                  for i in s.items],
        'events': [{'id': e.id, 'status': e.status, 'location': e.location,
                    'notes': e.notes, 'created_by': e.created_by,
                    'created_at': e.created_at.isoformat() if e.created_at else None}
                   for e in s.events]
    }), 200


@app.route('/api/shipments/', methods=['POST'])
@jwt_required()
def create_shipment():
    d = request.get_json()
    uid = int(get_jwt_identity())
    client = db.session.get(Client, d.get('client_id')) if d.get('client_id') else None
    s = Shipment(
        shipment_number=f"SHP-{datetime.now().year}-{random.randint(10000, 99999)}",
        organization_id=get_current_org(),
        client_id=client.id if client else None,
        client_name=client.company_name if client else d.get('clientName', 'Unknown'),
        origin=d.get('origin', ''), destination=d.get('destination', ''),
        carrier=d.get('carrier', ''), vessel_flight=d.get('vesselFlight', ''),
        status='Booked'
    )
    if d.get('etd'):
        s.etd = datetime.fromisoformat(d['etd'].replace('Z', ''))
    if d.get('eta'):
        s.eta = datetime.fromisoformat(d['eta'].replace('Z', ''))
    db.session.add(s)
    db.session.flush()
    db.session.add(TrackingEvent(shipment_id=s.id, status='Booked',
                                 location=s.origin, notes='Shipment created',
                                 created_by='system'))
    db.session.commit()
    log_activity('CREATE', 'shipment', s.id, s.shipment_number)
    notify(uid, '✅ Shipment Created', f'{s.shipment_number} created', 'success')
    socketio.emit('shipment_created', {'id': s.id, 'shipment_number': s.shipment_number}, broadcast=True)
    return jsonify({'msg': 'Created', 'id': s.id, 'shipment_number': s.shipment_number}), 201


@app.route('/api/shipments/<int:sid>', methods=['PUT'])
@jwt_required()
def update_shipment(sid):
    s = db.session.get(Shipment, sid)
    if not s:
        return jsonify({'msg': 'Not found'}), 404
    d = request.get_json()
    uid = int(get_jwt_identity())
    old_status = s.status
    for k in ['origin', 'destination', 'client_name', 'carrier',
              'vessel_flight', 'currency', 'payment_status']:
        if k in d:
            setattr(s, k, d[k])
    if 'etd' in d and d['etd']:
        s.etd = datetime.fromisoformat(d['etd'].replace('Z', ''))
    if 'eta' in d and d['eta']:
        s.eta = datetime.fromisoformat(d['eta'].replace('Z', ''))
    if 'status' in d:
        s.status = d['status']
        if d['status'] != old_status:
            db.session.add(TrackingEvent(shipment_id=s.id, status=d['status'],
                                         location=s.destination,
                                         notes=f'Status changed to {d["status"]}',
                                         created_by='system'))
    db.session.commit()
    log_activity('UPDATE', 'shipment', sid, f'Status: {s.status}')
    socketio.emit('shipment_updated', {'id': s.id, 'status': s.status,
                                        'client_name': s.client_name}, broadcast=True)
    return jsonify({'msg': 'Updated'}), 200


@app.route('/api/shipments/<int:sid>', methods=['DELETE'])
@jwt_required()
def delete_shipment(sid):
    s = db.session.get(Shipment, sid)
    if not s:
        return jsonify({'msg': 'Not found'}), 404
    s.is_deleted = True
    db.session.commit()
    log_activity('DELETE', 'shipment', sid, s.shipment_number)
    socketio.emit('shipment_deleted', {'id': sid}, broadcast=True)
    return jsonify({'msg': 'Deleted'}), 200


# ==================== ITEMS ====================
@app.route('/api/shipments/<int:sid>/items', methods=['POST'])
@jwt_required()
def add_item(sid):
    s = db.session.get(Shipment, sid)
    if not s:
        return jsonify({'msg': 'Not found'}), 404
    d = request.get_json()
    i = ShipmentItem(
        shipment_id=sid, description=d['description'],
        quantity=int(d.get('quantity', 1)),
        weight_kg=float(d.get('weight_kg', 0)),
        declared_value=float(d.get('declared_value', 0)),
        hs_code=d.get('hs_code', ''),
        package_type=d.get('package_type', '')
    )
    db.session.add(i)
    db.session.flush()
    s.total_value = sum(x.declared_value for x in s.items)
    db.session.commit()
    return jsonify({'msg': 'Item added', 'id': i.id}), 201


@app.route('/api/items/<int:iid>', methods=['DELETE'])
@jwt_required()
def delete_item(iid):
    i = db.session.get(ShipmentItem, iid)
    if not i:
        return jsonify({'msg': 'Not found'}), 404
    s = i.shipment
    db.session.delete(i)
    db.session.flush()
    s.total_value = sum(x.declared_value for x in s.items)
    db.session.commit()
    return jsonify({'msg': 'Removed'}), 200


# ==================== EVENTS ====================
@app.route('/api/shipments/<int:sid>/events', methods=['POST'])
@jwt_required()
def add_event(sid):
    s = db.session.get(Shipment, sid)
    if not s:
        return jsonify({'msg': 'Not found'}), 404
    d = request.get_json()
    uid = int(get_jwt_identity())
    u = db.session.get(User, uid)
    e = TrackingEvent(
        shipment_id=sid, status=d['status'],
        location=d.get('location', ''), notes=d.get('notes', ''),
        created_by=u.username if u else 'system'
    )
    db.session.add(e)
    if d['status'] != s.status:
        s.status = d['status']
    db.session.commit()
    try:
        if s.client_id:
            client = db.session.get(Client, s.client_id)
            if client and client.email:
                send_email(client.email,
                           f'📦 Shipment {s.shipment_number} - {e.status}',
                           shipment_email_html(s, e))
    except:
        pass
    socketio.emit('shipment_updated', {'id': s.id, 'status': s.status}, broadcast=True)
    return jsonify({'msg': 'Event added', 'id': e.id}), 201


# ==================== INVOICES ====================
@app.route('/api/invoices/', methods=['GET'])
@jwt_required()
def list_invoices():
    org_id = get_current_org()
    status = request.args.get('status')
    q = Invoice.query.filter_by(organization_id=org_id)
    if status and status != 'All':
        q = q.filter_by(status=status)
    invoices = q.order_by(Invoice.id.desc()).all()
    return jsonify([{
        'id': i.id, 'invoice_number': i.invoice_number,
        'shipment_id': i.shipment_id, 'client_name': i.client_name,
        'amount': i.amount, 'tax': i.tax, 'total': i.total,
        'paid_amount': i.paid_amount or 0, 'balance': i.balance,
        'currency': i.currency, 'status': i.status,
        'issue_date': i.issue_date.isoformat() if i.issue_date else None,
        'due_date': i.due_date.isoformat() if i.due_date else None,
        'paid_date': i.paid_date.isoformat() if i.paid_date else None,
        'notes': i.notes
    } for i in invoices]), 200


@app.route('/api/invoices/', methods=['POST'])
@jwt_required()
def create_invoice():
    d = request.get_json()
    s = db.session.get(Shipment, d.get('shipment_id'))
    if not s:
        return jsonify({'msg': 'Shipment not found'}), 404
    amount = float(d.get('amount', s.total_value))
    tax_rate = float(d.get('tax_rate', 0.18))
    tax = amount * tax_rate
    inv = Invoice(
        invoice_number=f"INV-{datetime.now().year}-{random.randint(10000, 99999)}",
        organization_id=get_current_org(),
        shipment_id=s.id, client_name=s.client_name,
        amount=amount, tax=tax, total=amount + tax,
        currency=s.currency,
        due_date=datetime.utcnow() + timedelta(days=30),
        notes=d.get('notes', '')
    )
    s.payment_status = 'Invoiced'
    db.session.add(inv)
    db.session.commit()
    return jsonify({'msg': 'Invoice created', 'id': inv.id,
                    'invoice_number': inv.invoice_number}), 201


@app.route('/api/invoices/<int:iid>/mark-paid', methods=['PUT'])
@jwt_required()
def mark_paid(iid):
    inv = db.session.get(Invoice, iid)
    if not inv:
        return jsonify({'msg': 'Not found'}), 404
    inv.paid_amount = inv.total
    inv.status = 'Paid'
    inv.paid_date = datetime.utcnow()
    if inv.shipment_id:
        s = db.session.get(Shipment, inv.shipment_id)
        if s:
            s.payment_status = 'Paid'
    db.session.commit()
    return jsonify({'msg': 'Marked paid'}), 200


@app.route('/api/invoices/<int:iid>', methods=['DELETE'])
@jwt_required()
def delete_invoice(iid):
    inv = db.session.get(Invoice, iid)
    if not inv:
        return jsonify({'msg': 'Not found'}), 404
    db.session.delete(inv)
    db.session.commit()
    return jsonify({'msg': 'Deleted'}), 200


# ==================== PAYMENTS ====================
@app.route('/api/invoices/<int:iid>/payments', methods=['POST'])
@jwt_required()
def record_payment(iid):
    inv = db.session.get(Invoice, iid)
    if not inv:
        return jsonify({'msg': 'Not found'}), 404
    d = request.get_json()
    amount = float(d.get('amount', 0))
    if amount <= 0:
        return jsonify({'msg': 'Amount must be > 0'}), 400
    uid = int(get_jwt_identity())
    u = db.session.get(User, uid)
    p = Payment(
        invoice_id=iid, amount=amount,
        method=d.get('method', 'Bank Transfer'),
        reference=d.get('reference', ''), notes=d.get('notes', ''),
        recorded_by=u.username if u else 'system'
    )
    db.session.add(p)
    inv.paid_amount = (inv.paid_amount or 0) + amount
    if inv.paid_amount >= inv.total:
        inv.status = 'Paid'
        inv.paid_date = datetime.utcnow()
    elif inv.paid_amount > 0:
        inv.status = 'Partial'
    if inv.shipment_id:
        s = db.session.get(Shipment, inv.shipment_id)
        if s:
            s.payment_status = inv.status
    db.session.commit()
    return jsonify({'msg': 'Payment recorded', 'balance': inv.balance, 'status': inv.status}), 201


@app.route('/api/invoices/<int:iid>/payments', methods=['GET'])
@jwt_required()
def list_payments(iid):
    payments = Payment.query.filter_by(invoice_id=iid).order_by(Payment.id.desc()).all()
    return jsonify([{
        'id': p.id, 'amount': p.amount, 'method': p.method,
        'reference': p.reference, 'notes': p.notes,
        'paid_date': p.paid_date.isoformat() if p.paid_date else None,
        'recorded_by': p.recorded_by
    } for p in payments]), 200


@app.route('/api/payments', methods=['GET'])
@jwt_required()
def all_payments():
    payments = Payment.query.order_by(Payment.id.desc()).limit(100).all()
    result = []
    for p in payments:
        inv = db.session.get(Invoice, p.invoice_id)
        result.append({
            'id': p.id, 'invoice_id': p.invoice_id,
            'invoice_number': inv.invoice_number if inv else '—',
            'client_name': inv.client_name if inv else '—',
            'amount': p.amount, 'method': p.method,
            'reference': p.reference,
            'paid_date': p.paid_date.isoformat() if p.paid_date else None,
            'recorded_by': p.recorded_by
        })
    return jsonify(result), 200


# ==================== DOCUMENTS ====================
@app.route('/api/shipments/<int:sid>/documents', methods=['POST'])
@jwt_required()
def upload_document(sid):
    s = db.session.get(Shipment, sid)
    if not s:
        return jsonify({'msg': 'Not found'}), 404
    if 'file' not in request.files:
        return jsonify({'msg': 'No file'}), 400
    file = request.files['file']
    if not file or not file.filename:
        return jsonify({'msg': 'Empty file'}), 400
    if not allowed_file(file.filename):
        return jsonify({'msg': 'File type not allowed'}), 400
    orig = file.filename
    safe = secure_filename(f"{sid}_{int(datetime.utcnow().timestamp())}_{orig}")
    path = os.path.join(app.config['UPLOAD_FOLDER'], safe)
    file.save(path)
    size = os.path.getsize(path)
    mime = file.content_type or mimetypes.guess_type(orig)[0] or 'application/octet-stream'
    uid = int(get_jwt_identity())
    u = db.session.get(User, uid)
    doc = Document(
        shipment_id=sid, filename=safe, original_name=orig,
        file_type=request.form.get('file_type', 'General'),
        file_size=size, mime_type=mime,
        uploaded_by=u.username if u else 'system'
    )
    db.session.add(doc)
    db.session.commit()
    return jsonify({'msg': 'Uploaded', 'id': doc.id}), 201


@app.route('/api/shipments/<int:sid>/documents', methods=['GET'])
@jwt_required()
def list_documents(sid):
    docs = Document.query.filter_by(shipment_id=sid).order_by(Document.id.desc()).all()
    return jsonify([{
        'id': d.id, 'original_name': d.original_name, 'file_type': d.file_type,
        'file_size': d.file_size, 'mime_type': d.mime_type,
        'uploaded_by': d.uploaded_by,
        'uploaded_at': d.uploaded_at.isoformat() if d.uploaded_at else None
    } for d in docs]), 200


@app.route('/api/documents/<int:did>/download', methods=['GET'])
@jwt_required()
def download_document(did):
    d = db.session.get(Document, did)
    if not d:
        return jsonify({'msg': 'Not found'}), 404
    path = os.path.join(app.config['UPLOAD_FOLDER'], d.filename)
    if not os.path.exists(path):
        return jsonify({'msg': 'File missing'}), 404
    return send_file(path, as_attachment=True, download_name=d.original_name)


@app.route('/api/documents/<int:did>', methods=['DELETE'])
@jwt_required()
def delete_document(did):
    d = db.session.get(Document, did)
    if not d:
        return jsonify({'msg': 'Not found'}), 404
    path = os.path.join(app.config['UPLOAD_FOLDER'], d.filename)
    try:
        if os.path.exists(path):
            os.remove(path)
    except:
        pass
    db.session.delete(d)
    db.session.commit()
    return jsonify({'msg': 'Deleted'}), 200


# ==================== DASHBOARD ====================
@app.route('/api/dashboard/stats', methods=['GET'])
@jwt_required()
def dashboard_stats():
    org_id = get_current_org()
    base = Shipment.query.filter_by(is_deleted=False, organization_id=org_id)
    total = base.count()
    in_transit = base.filter_by(status='In Transit').count()
    delivered = base.filter_by(status='Delivered').count()
    confirmed = base.filter_by(status='Confirmed').count()
    booked = base.filter_by(status='Booked').count()
    revenue = db.session.query(func.coalesce(func.sum(Shipment.total_value), 0))\
        .filter_by(is_deleted=False, organization_id=org_id).scalar()
    paid_revenue = db.session.query(func.coalesce(func.sum(Shipment.total_value), 0))\
        .filter_by(is_deleted=False, organization_id=org_id, payment_status='Paid').scalar()
    return jsonify({
        'total': total, 'inTransit': in_transit, 'delivered': delivered,
        'revenue': float(revenue), 'paidRevenue': float(paid_revenue),
        'pendingRevenue': float(revenue - paid_revenue),
        'labels': ['Booked', 'Confirmed', 'In Transit', 'Delivered'],
        'values': [booked, confirmed, in_transit, delivered]
    }), 200


@app.route('/api/dashboard/trends', methods=['GET'])
@jwt_required()
def trends():
    org_id = get_current_org()
    t = db.session.query(func.date(Shipment.created_at), func.count(Shipment.id))\
        .filter(Shipment.is_deleted == False, Shipment.organization_id == org_id)\
        .group_by(func.date(Shipment.created_at))\
        .order_by(func.date(Shipment.created_at)).all()
    return jsonify([{'date': str(r[0]), 'count': r[1]} for r in t]), 200


@app.route('/api/dashboard/advanced-stats', methods=['GET'])
@jwt_required()
def adv_stats():
    org_id = get_current_org()
    top_clients = db.session.query(Shipment.client_name, func.count(Shipment.id))\
        .filter(Shipment.is_deleted == False, Shipment.organization_id == org_id)\
        .group_by(Shipment.client_name)\
        .order_by(func.count(Shipment.id).desc()).limit(5).all()
    top_routes = db.session.query(Shipment.origin, Shipment.destination, func.count(Shipment.id))\
        .filter(Shipment.is_deleted == False, Shipment.organization_id == org_id)\
        .group_by(Shipment.origin, Shipment.destination)\
        .order_by(func.count(Shipment.id).desc()).limit(5).all()
    return jsonify({
        'topClients': {
            'labels': [c[0] or 'N/A' for c in top_clients],
            'values': [c[1] for c in top_clients]
        },
        'topRoutes': {
            'labels': [f"{r[0]} → {r[1]}" for r in top_routes],
            'values': [r[2] for r in top_routes]
        }
    }), 200


@app.route('/api/dashboard/revenue', methods=['GET'])
@jwt_required()
def revenue_analytics():
    org_id = get_current_org()
    invoices = Invoice.query.filter_by(organization_id=org_id).all()
    ti = sum(i.total or 0 for i in invoices)
    tp = sum(i.paid_amount or 0 for i in invoices)
    monthly = db.session.query(
        func.to_char(Invoice.issue_date, 'YYYY-MM').label('month'),
        func.coalesce(func.sum(Invoice.total), 0)
    ).filter(Invoice.organization_id == org_id)\
     .group_by(func.to_char(Invoice.issue_date, 'YYYY-MM'))\
     .order_by(func.to_char(Invoice.issue_date, 'YYYY-MM').desc()).limit(6).all()
    monthly.reverse()
    return jsonify({
        'totalInvoiced': round(ti, 2), 'totalPaid': round(tp, 2),
        'outstanding': round(ti - tp, 2),
        'months': [m[0] for m in monthly],
        'revenueByMonth': [float(m[1]) for m in monthly]
    }), 200


@app.route('/api/dashboard/carrier-performance', methods=['GET'])
@jwt_required()
def carrier_performance():
    org_id = get_current_org()
    rows = db.session.query(
        Shipment.carrier,
        func.count(Shipment.id),
        func.coalesce(func.sum(Shipment.total_value), 0),
        func.sum(func.case((Shipment.status == 'Delivered', 1), else_=0))
    ).filter(Shipment.is_deleted == False, Shipment.organization_id == org_id,
             Shipment.carrier.isnot(None))\
     .group_by(Shipment.carrier)\
     .order_by(func.count(Shipment.id).desc()).limit(10).all()
    return jsonify([{
        'carrier': r[0] or 'Unknown', 'shipments': r[1],
        'revenue': float(r[2]), 'delivered': int(r[3] or 0),
        'delivery_rate': round((int(r[3] or 0) / r[1] * 100), 1) if r[1] else 0
    } for r in rows]), 200


@app.route('/api/dashboard/heatmap', methods=['GET'])
@jwt_required()
def geographic_heatmap():
    org_id = get_current_org()
    rows = db.session.query(
        Shipment.origin, Shipment.latitude, Shipment.longitude,
        func.count(Shipment.id), func.coalesce(func.sum(Shipment.total_value), 0)
    ).filter(Shipment.is_deleted == False, Shipment.organization_id == org_id,
             Shipment.latitude.isnot(None))\
     .group_by(Shipment.origin, Shipment.latitude, Shipment.longitude).all()
    return jsonify([{
        'city': r[0], 'lat': r[1], 'lng': r[2],
        'count': r[3], 'revenue': float(r[4])
    } for r in rows]), 200


@app.route('/api/dashboard/profitability', methods=['GET'])
@jwt_required()
def profitability():
    org_id = get_current_org()
    rows = db.session.query(
        Shipment.client_name,
        func.coalesce(func.sum(Shipment.total_value), 0),
        func.count(Shipment.id)
    ).filter(Shipment.is_deleted == False, Shipment.organization_id == org_id)\
     .group_by(Shipment.client_name)\
     .order_by(func.sum(Shipment.total_value).desc()).limit(5).all()
    return jsonify([{
        'client': r[0] or 'Unknown', 'revenue': float(r[1]), 'shipments': r[2]
    } for r in rows]), 200


# ==================== AI DELAY RISK ====================
@app.route('/api/shipments/<int:sid>/delay-risk', methods=['GET'])
@jwt_required()
def delay_risk(sid):
    s = db.session.get(Shipment, sid)
    if not s:
        return jsonify({'msg': 'Not found'}), 404
    score = 0
    factors = []
    if s.status == 'Booked':
        score += 30
        factors.append('Still in booking stage')
    elif s.status == 'In Transit':
        score += 15
        factors.append('Currently in transit')
    long_routes = [('Shanghai', 'New York'), ('Tokyo', 'London'), ('Sydney', 'Moscow')]
    if (s.origin, s.destination) in long_routes:
        score += 25
        factors.append('Long-distance route')
    if s.created_at:
        age = (datetime.utcnow() - s.created_at).days
        if age > 7:
            score += 20
            factors.append(f'{age} days old')
    if not s.carrier:
        score += 10
        factors.append('No carrier assigned')
    if s.payment_status == 'Unpaid':
        score += 15
        factors.append('Payment pending')
    score = min(score, 100)
    level = 'LOW' if score < 30 else 'MEDIUM' if score < 60 else 'HIGH'
    return jsonify({'risk_score': score, 'risk_level': level,
                    'factors': factors or ['No risk factors']}), 200


# ==================== ACTIVITY LOG ====================
@app.route('/api/activity-logs', methods=['GET'])
@jwt_required()
def activity_logs():
    logs = ActivityLog.query.order_by(ActivityLog.id.desc()).limit(200).all()
    return jsonify([{
        'id': l.id, 'username': l.username, 'action': l.action,
        'entity': l.entity, 'entity_id': l.entity_id,
        'details': l.details, 'ip_address': l.ip_address,
        'created_at': l.created_at.isoformat() if l.created_at else None
    } for l in logs]), 200


# ==================== NOTIFICATIONS ====================
@app.route('/api/notifications/', methods=['GET'])
@jwt_required()
def list_notifs():
    uid = int(get_jwt_identity())
    notifs = Notification.query.filter_by(user_id=uid)\
        .order_by(Notification.id.desc()).limit(50).all()
    unread = Notification.query.filter_by(user_id=uid, is_read=False).count()
    return jsonify({
        'unread': unread,
        'items': [{
            'id': n.id, 'title': n.title, 'message': n.message,
            'type': n.type, 'is_read': n.is_read, 'link': n.link,
            'created_at': n.created_at.isoformat() if n.created_at else None
        } for n in notifs]
    }), 200


@app.route('/api/notifications/<int:nid>/read', methods=['PUT'])
@jwt_required()
def mark_notif_read(nid):
    n = db.session.get(Notification, nid)
    if not n:
        return jsonify({'msg': 'Not found'}), 404
    n.is_read = True
    db.session.commit()
    return jsonify({'msg': 'Read'}), 200


@app.route('/api/notifications/read-all', methods=['PUT'])
@jwt_required()
def mark_all_read():
    uid = int(get_jwt_identity())
    Notification.query.filter_by(user_id=uid, is_read=False).update({'is_read': True})
    db.session.commit()
    return jsonify({'msg': 'All read'}), 200


# ==================== ADMIN PANEL ====================
@app.route('/api/admin/stats', methods=['GET'])
@jwt_required()
@admin_required()
def admin_stats():
    week_ago = datetime.utcnow() - timedelta(days=7)
    return jsonify({
        'totalUsers': User.query.count(),
        'activeUsers': User.query.filter_by(is_active=True).count(),
        'totalOrganizations': Organization.query.count(),
        'activeOrganizations': Organization.query.filter_by(is_active=True).count(),
        'totalShipments': Shipment.query.filter_by(is_deleted=False).count(),
        'totalClients': Client.query.count(),
        'totalInvoices': Invoice.query.count(),
        'totalRevenue': float(db.session.query(func.coalesce(func.sum(Invoice.total), 0)).scalar()),
        'recentLogins': ActivityLog.query.filter(
            ActivityLog.action == 'LOGIN',
            ActivityLog.created_at >= week_ago
        ).count()
    }), 200


@app.route('/api/admin/users', methods=['GET'])
@jwt_required()
@admin_required()
def admin_list_users():
    users = User.query.order_by(User.id.desc()).all()
    result = []
    for u in users:
        org = db.session.get(Organization, u.organization_id) if u.organization_id else None
        result.append({
            'id': u.id, 'username': u.username, 'email': u.email,
            'role': u.role, 'is_active': u.is_active,
            'department': u.department,
            'organization': org.name if org else 'None',
            'organization_id': u.organization_id,
            'last_login': u.last_login.isoformat() if u.last_login else None,
            'created_at': u.created_at.isoformat() if u.created_at else None
        })
    return jsonify(result), 200


@app.route('/api/admin/users/<int:uid>', methods=['PUT'])
@jwt_required()
@admin_required()
def admin_update_user(uid):
    u = db.session.get(User, uid)
    if not u:
        return jsonify({'msg': 'User not found'}), 404
    d = request.get_json()
    if 'role' in d and d['role'] in ['admin', 'user']:
        u.role = d['role']
    if 'is_active' in d:
        u.is_active = bool(d['is_active'])
    if 'email' in d:
        u.email = d['email']
    if 'department' in d:
        u.department = d['department']
    if 'password' in d and d['password']:
        u.set_password(d['password'])
    db.session.commit()
    log_activity('UPDATE', 'user', uid, f'Admin updated: {u.username}')
    return jsonify({'msg': 'User updated'}), 200


@app.route('/api/admin/users/<int:uid>', methods=['DELETE'])
@jwt_required()
@admin_required()
def admin_delete_user(uid):
    if uid == int(get_jwt_identity()):
        return jsonify({'msg': 'Cannot delete yourself'}), 400
    u = db.session.get(User, uid)
    if not u:
        return jsonify({'msg': 'Not found'}), 404
    db.session.delete(u)
    db.session.commit()
    return jsonify({'msg': 'Deleted'}), 200


@app.route('/api/admin/organizations', methods=['GET'])
@jwt_required()
@admin_required()
def admin_list_orgs():
    orgs = Organization.query.order_by(Organization.id.desc()).all()
    result = []
    for o in orgs:
        result.append({
            'id': o.id, 'name': o.name, 'slug': o.slug,
            'plan': o.plan, 'is_active': o.is_active,
            'users': User.query.filter_by(organization_id=o.id).count(),
            'shipments': Shipment.query.filter_by(organization_id=o.id, is_deleted=False).count(),
            'created_at': o.created_at.isoformat() if o.created_at else None
        })
    return jsonify(result), 200


@app.route('/api/admin/organizations/<int:oid>', methods=['PUT'])
@jwt_required()
@admin_required()
def admin_update_org(oid):
    o = db.session.get(Organization, oid)
    if not o:
        return jsonify({'msg': 'Not found'}), 404
    d = request.get_json()
    if 'name' in d:
        o.name = d['name']
    if 'plan' in d:
        o.plan = d['plan']
    if 'is_active' in d:
        o.is_active = bool(d['is_active'])
    db.session.commit()
    return jsonify({'msg': 'Updated'}), 200


# ==================== REPORTS ====================
@app.route('/api/reports/excel', methods=['GET'])
@jwt_required()
def export_excel():
    from openpyxl import Workbook
    from io import BytesIO
    org_id = get_current_org()
    shipments = Shipment.query.filter_by(is_deleted=False, organization_id=org_id).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Shipments"
    ws.append(['ID', 'Shipment #', 'Client', 'Origin', 'Destination',
               'Status', 'Carrier', 'Total Value', 'Payment'])
    for s in shipments:
        ws.append([s.id, s.shipment_number, s.client_name, s.origin,
                   s.destination, s.status, s.carrier,
                   s.total_value, s.payment_status])
    for col in ws.columns:
        max_length = max((len(str(cell.value)) if cell.value else 0) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_length + 2
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output, as_attachment=True,
        download_name='FreightPulse_Shipments.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


@app.route('/api/reports/pdf', methods=['GET'])
@jwt_required()
def export_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from io import BytesIO
    org_id = get_current_org()
    shipments = Shipment.query.filter_by(is_deleted=False, organization_id=org_id).all()
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=30, leftMargin=30,
                            topMargin=30, bottomMargin=18)
    elements = []
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('T', parent=styles['Heading1'], fontSize=18,
                                  textColor=colors.HexColor('#0a1628'),
                                  alignment=1, spaceAfter=12)
    elements.append(Paragraph("FreightPulse - Shipment Report", title_style))
    elements.append(Spacer(1, 0.2 * inch))
    data = [['#', 'Shipment #', 'Client', 'Origin', 'Destination', 'Status']]
    for idx, s in enumerate(shipments, start=1):
        data.append([str(idx), s.shipment_number, s.client_name or 'N/A',
                     s.origin or 'N/A', s.destination or 'N/A', s.status or 'Draft'])
    table = Table(data, colWidths=[0.5 * inch, 1.4 * inch, 1.5 * inch,
                                   1.3 * inch, 1.3 * inch, 1.2 * inch])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#C9A84C')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    elements.append(table)
    doc.build(elements)
    buffer.seek(0)
    return send_file(buffer, as_attachment=True,
                     download_name='FreightPulse_Report.pdf',
                     mimetype='application/pdf')


# ==================== SEED ====================
@app.route('/api/seed', methods=['POST', 'OPTIONS'])
@jwt_required()
def seed():
    if request.method == 'OPTIONS':
        return '', 200
    uid = int(get_jwt_identity())
    org_id = get_current_org()

    Shipment.query.filter_by(organization_id=org_id).delete()
    Client.query.filter_by(organization_id=org_id).delete()
    Invoice.query.filter_by(organization_id=org_id).delete()
    db.session.commit()

    cities = [
        ("Shanghai", 31.23, 121.47), ("New York", 40.71, -74.00),
        ("London", 51.51, -0.13), ("Tokyo", 35.68, 139.65),
        ("Mumbai", 19.08, 72.88), ("Singapore", 1.35, 103.82),
        ("Dubai", 25.20, 55.27), ("Sydney", -33.87, 151.21),
        ("Toronto", 43.65, -79.38), ("Los Angeles", 34.05, -118.24),
        ("Rio de Janeiro", -22.91, -43.17), ("Cape Town", -33.92, 18.42)
    ]
    statuses = ['Booked', 'Confirmed', 'In Transit', 'Delivered']
    currencies = ['USD', 'EUR', 'INR', 'GBP']

    clients = []
    for _ in range(8):
        c = Client(
            organization_id=org_id,
            company_name=fake.company(),
            contact_person=fake.name(),
            email=fake.email(), phone=fake.phone_number()[:20],
            city=fake.city(), country=fake.country(),
            tax_id=fake.bothify('??####??')
        )
        db.session.add(c)
        clients.append(c)
    db.session.flush()

    for _ in range(30):
        o = random.choice(cities)
        d_ = random.choice(cities)
        cl = random.choice(clients)
        status = random.choice(statuses)
        s = Shipment(
            organization_id=org_id,
            shipment_number=f"SHP-{fake.year()}-{fake.random_number(digits=5)}",
            client_id=cl.id, client_name=cl.company_name,
            origin=o[0], destination=d_[0], status=status,
            carrier=fake.company(), currency=random.choice(currencies),
            latitude=o[1], longitude=o[2],
            etd=datetime.utcnow() + timedelta(days=random.randint(1, 10)),
            eta=datetime.utcnow() + timedelta(days=random.randint(11, 30))
        )
        db.session.add(s)
        db.session.flush()
        for _ in range(random.randint(1, 3)):
            db.session.add(ShipmentItem(
                shipment_id=s.id, description=fake.catch_phrase(),
                quantity=random.randint(1, 100),
                weight_kg=random.uniform(10, 500),
                declared_value=random.uniform(100, 5000),
                hs_code=fake.bothify('####.##'),
                package_type=random.choice(['Cartons', 'Pallets', 'Drums', 'Bags'])
            ))
        db.session.flush()
        s.total_value = sum(i.declared_value for i in s.items)
        for st in ['Booked', 'Confirmed', 'In Transit', 'Delivered'][:statuses.index(status) + 1]:
            db.session.add(TrackingEvent(
                shipment_id=s.id, status=st,
                location=o[0] if st in ['Booked', 'Confirmed'] else d_[0],
                notes=f'{st} event', created_by='system'
            ))
        s.payment_status = random.choice(['Unpaid', 'Invoiced', 'Paid'])

    db.session.commit()
    socketio.emit('data_refresh', {}, broadcast=True)
    return jsonify({'msg': '30 shipments seeded!'}), 201


@app.route('/api/clear', methods=['DELETE', 'OPTIONS'])
@jwt_required()
@admin_required()
def clear():
    if request.method == 'OPTIONS':
        return '', 200
    org_id = get_current_org()
    Shipment.query.filter_by(organization_id=org_id).delete()
    Client.query.filter_by(organization_id=org_id).delete()
    Invoice.query.filter_by(organization_id=org_id).delete()
    db.session.commit()
    socketio.emit('data_refresh', {}, broadcast=True)
    return jsonify({'msg': 'Data cleared'}), 200


# ==================== SOCKETIO ====================
@socketio.on('connect')
def on_conn():
    print('🔌 Client connected')


@socketio.on('disconnect')
def on_disconnect():
    print('🔌 Client disconnected')


# ==================== RUN ====================
if __name__ == '__main__':
    socketio.run(app, debug=True, port=5000)